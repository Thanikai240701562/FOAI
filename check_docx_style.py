import docx

doc = docx.Document(r'c:\AI & Web project\115-Traffic and Weather Data Analysis for Smart Travel Suggestion.docx')
s = doc.sections[0]
print(f"Page dimensions: {s.page_width.pt} x {s.page_height.pt} pt")
print(f"Margins (L, R, T, B): {s.left_margin.pt}, {s.right_margin.pt}, {s.top_margin.pt}, {s.bottom_margin.pt} pt")

indices = [0, 1, 2, 3, 4, 5, 6, 10, 16, 17, 37, 38, 70, 96, 98, 102, 104, 105]
for idx in indices:
    if idx < len(doc.paragraphs):
        p = doc.paragraphs[idx]
        r = p.runs[0] if p.runs else None
        print(f"P{idx}: text='{p.text[:40]}' | style='{p.style.name}' | align={p.alignment} | space_before={p.paragraph_format.space_before} | space_after={p.paragraph_format.space_after} | font={r.font.name if r else None} | size={r.font.size.pt if r and r.font.size else None} | bold={r.font.bold if r else None}")
